<<<<<<< Updated upstream
from typing import Dict, List, Union
=======
from typing import Dict, List, Union, Iterator, Tuple
from io import StringIO, BytesIO
import csv
import json
import os
import tempfile
from openpyxl import Workbook
>>>>>>> Stashed changes
from .formatters import BaseFormatter, CSVFormatter, JSONFormatter, XLSXFormatter


class DataExporter:
    """
    Export data using DRF serializers with built-in format support.

    Usage:
        # New simplified interface
        exporter = DataExporter(BookSerializer, format_type='csv')
        filename, content = exporter.export('books_export', queryset)

        # Legacy interface (still supported)
        exporter = DataExporter(BookSerializer)
        csv_string = exporter.to_string(queryset, CSVFormatter())
    """

    # Available formatters
    FORMATTERS = {
        "csv": CSVFormatter,
        "json": JSONFormatter,
        "xlsx": XLSXFormatter,
    }

    def __init__(self, serializer_class, format_type: str = None, **serializer_kwargs):
        """
        Initialize exporter with serializer and optional format type.

        Args:
            serializer_class: DRF serializer class to use for data serialization
            format_type: Optional format type (csv, json, xlsx). If provided, enables export() method.
            **serializer_kwargs: Additional kwargs to pass to serializer
        """
        self.serializer_class = serializer_class
        self.serializer_kwargs = serializer_kwargs
        self.format_type = format_type
        self.formatter = None

        if format_type:
            if format_type not in self.FORMATTERS:
                raise ValueError(f"Unsupported format: {format_type}. Available: {list(self.FORMATTERS.keys())}")
            # Create formatter with default options
            self.formatter = self._create_formatter(format_type)

    def _create_formatter(self, format_type: str) -> BaseFormatter:
        """Create formatter instance with appropriate options."""
        formatter_class = self.FORMATTERS[format_type]

        # Apply format-specific options
        if format_type == "xlsx":
            return formatter_class(list_joiner=", ")
        else:
            return formatter_class()

    def serialize(self, queryset) -> List[Dict]:
        """QuerySet → list of dicts"""
        serializer = self.serializer_class(
            queryset,
            many=True,
            **self.serializer_kwargs
        )
        return serializer.data

    def export(self, filename: str, queryset) -> tuple[str, Union[str, bytes]]:
        """
        Export queryset to file with configured format.

        Args:
            filename: Base filename (without extension)
            queryset: Django QuerySet to export

        Returns:
            Tuple of (filename_with_extension, content)

        Raises:
            ValueError: If format_type was not provided during initialization
        """
        if not self.formatter:
            raise ValueError("format_type must be provided during initialization to use export() method")

        data = self.serialize(queryset)
        content = self.formatter.encode(data)
        full_filename = f"{filename}.{self.formatter.extension}"

        return full_filename, content

    def to_string(self, queryset, formatter: BaseFormatter) -> Union[str, bytes]:
        """Export to formatted string (legacy interface)"""
        data = self.serialize(queryset)
        return formatter.encode(data)

    def to_file(self, queryset, filepath: str, formatter: BaseFormatter) -> str:
        """Export to file (legacy interface)"""
        content = self.to_string(queryset, formatter)
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(content)
        return filepath

    @classmethod
    def get_available_formats(cls) -> List[str]:
        """Get list of available export formats."""
        return list(cls.FORMATTERS.keys())
<<<<<<< Updated upstream
=======

    def _iter_chunks(self, queryset, chunk_size: int) -> Iterator[List[Dict]]:
        """
        Generator that yields serialized chunks of the queryset.

        Args:
            queryset: Django QuerySet to iterate
            chunk_size: Number of records per chunk

        Yields:
            List of serialized dictionaries for each chunk
        """
        total_count = queryset.count()
        for offset in range(0, total_count, chunk_size):
            chunk = queryset[offset:offset + chunk_size]
            serialized_chunk = self.serialize(chunk)
            if serialized_chunk:
                yield serialized_chunk

    def _extract_fieldnames(self, serialized_data: List[Dict]) -> List[str]:
        """
        Extract unique field names from serialized data in order of appearance.

        Args:
            serialized_data: List of dictionaries from serializer

        Returns:
            List of unique field names
        """
        fieldnames = []
        for row in serialized_data:
            for key in row.keys():
                if key not in fieldnames:
                    fieldnames.append(key)
        return fieldnames

    def _build_filename(self, base_filename: str) -> str:
        """Build full filename with extension based on format type."""
        return f"{base_filename}.{self.formatter.extension}"

    def export_chunked(self, filename: str, queryset, chunk_size: int = 1000) -> Tuple[str, Union[str, bytes]]:
        """
        Export queryset in chunks to avoid memory issues with large datasets.

        Args:
            filename: Base filename (without extension)
            queryset: Django QuerySet to export
            chunk_size: Number of records to process at a time

        Returns:
            Tuple of (filename_with_extension, content)
        """
        if not self.formatter:
            raise ValueError("format_type must be provided during initialization to use export_chunked() method")

        # Route to appropriate chunked export method based on format
        if self.format_type == "csv":
            return self._export_chunked_csv(filename, queryset, chunk_size)
        elif self.format_type == "json":
            return self._export_chunked_json(filename, queryset, chunk_size)
        elif self.format_type == "xlsx":
            return self._export_chunked_xlsx(filename, queryset, chunk_size)
        else:
            # Fallback to regular export for unknown formats
            return self.export(filename, queryset)

    def _export_chunked_csv(self, filename: str, queryset, chunk_size: int) -> Tuple[str, str]:
        """Export to CSV in chunks."""
        output = StringIO()
        writer = None
        fieldnames = None

        for serialized_chunk in self._iter_chunks(queryset, chunk_size):
            # Flatten data for CSV if needed
            if isinstance(self.formatter, CSVFormatter) and self.formatter.flatten:
                serialized_chunk = [self.formatter._flatten(row) for row in serialized_chunk]

            # Initialize writer with fieldnames from first chunk
            if writer is None:
                fieldnames = self._extract_fieldnames(serialized_chunk)

                # Write header
                if isinstance(self.formatter, CSVFormatter) and self.formatter.prettify_headers:
                    header_map = {key: self.formatter._prettify_header(key) for key in fieldnames}
                    pretty_headers = [header_map[key] for key in fieldnames]
                    writer = csv.writer(output, delimiter=self.formatter.delimiter)
                    writer.writerow(pretty_headers)
                else:
                    writer = csv.DictWriter(output, fieldnames=fieldnames, delimiter=self.formatter.delimiter)
                    writer.writeheader()

            # Write data rows
            for row in serialized_chunk:
                if isinstance(self.formatter, CSVFormatter) and self.formatter.prettify_headers:
                    writer.writerow([row.get(key, "") for key in fieldnames])
                else:
                    writer.writerow(row)

        return self._build_filename(filename), output.getvalue()

    def _export_chunked_json(self, filename: str, queryset, chunk_size: int) -> Tuple[str, str]:
        """Export to JSON in chunks."""
        output = StringIO()
        output.write("[")
        first_item = True

        for serialized_chunk in self._iter_chunks(queryset, chunk_size):
            for item in serialized_chunk:
                if not first_item:
                    output.write(",")
                output.write("\n  ")
                output.write(json.dumps(item, indent=None, default=str))
                first_item = False

        output.write("\n]")
        return self._build_filename(filename), output.getvalue()

    def _export_chunked_xlsx(self, filename: str, queryset, chunk_size: int) -> Tuple[str, bytes]:
        """Export to XLSX in chunks."""
        wb = Workbook()
        ws = wb.active
        fieldnames = None

        for serialized_chunk in self._iter_chunks(queryset, chunk_size):
            # Initialize headers from first chunk
            if fieldnames is None:
                fieldnames = self._extract_fieldnames(serialized_chunk)

                # Write header row
                if isinstance(self.formatter, XLSXFormatter) and self.formatter.prettify_headers:
                    headers = [self.formatter._prettify_header(key) for key in fieldnames]
                else:
                    headers = fieldnames
                ws.append(headers)

            # Write data rows
            for row in serialized_chunk:
                formatted_row = []
                for key in fieldnames:
                    value = row.get(key, "")
                    if isinstance(self.formatter, XLSXFormatter):
                        value = self.formatter._format_value(value)
                    formatted_row.append(value)
                ws.append(formatted_row)

        # Save to bytes
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return self._build_filename(filename), output.getvalue()

    def export_chunked_to_file(
        self, filename: str, queryset, chunk_size: int = 1000
    ) -> Tuple[str, str]:
        """
        Export queryset in chunks directly to a temporary file on disk.
        This method is memory-efficient for very large exports (>5000 items).

        Args:
            filename: Base filename (without extension)
            queryset: Django QuerySet to export
            chunk_size: Number of records to process at a time

        Returns:
            Tuple of (filename_with_extension, temp_file_path)
            The caller is responsible for cleaning up the temp file.
        """
        if not self.formatter:
            raise ValueError("format_type must be provided during initialization")

        # Route to appropriate file-based export method
        if self.format_type == "csv":
            return self._export_chunked_csv_to_file(filename, queryset, chunk_size)
        elif self.format_type == "json":
            return self._export_chunked_json_to_file(filename, queryset, chunk_size)
        elif self.format_type == "xlsx":
            return self._export_chunked_xlsx_to_file(filename, queryset, chunk_size)
        else:
            raise ValueError(f"Unsupported format for file-based export: {self.format_type}")

    def _export_chunked_csv_to_file(
        self, filename: str, queryset, chunk_size: int
    ) -> Tuple[str, str]:
        """Export to CSV in chunks, writing directly to a temp file."""
        temp_file = tempfile.NamedTemporaryFile(
            mode='w', delete=False, suffix='.csv', encoding='utf-8', newline=''
        )
        try:
            writer = None
            fieldnames = None

            for serialized_chunk in self._iter_chunks(queryset, chunk_size):
                # Flatten data for CSV if needed
                if isinstance(self.formatter, CSVFormatter) and self.formatter.flatten:
                    serialized_chunk = [self.formatter._flatten(row) for row in serialized_chunk]

                # Initialize writer with fieldnames from first chunk
                if writer is None:
                    fieldnames = self._extract_fieldnames(serialized_chunk)

                    # Write header
                    if isinstance(self.formatter, CSVFormatter) and self.formatter.prettify_headers:
                        header_map = {key: self.formatter._prettify_header(key) for key in fieldnames}
                        pretty_headers = [header_map[key] for key in fieldnames]
                        writer = csv.writer(temp_file, delimiter=self.formatter.delimiter)
                        writer.writerow(pretty_headers)
                    else:
                        writer = csv.DictWriter(
                            temp_file, fieldnames=fieldnames, delimiter=self.formatter.delimiter
                        )
                        writer.writeheader()

                # Write data rows
                for row in serialized_chunk:
                    if isinstance(self.formatter, CSVFormatter) and self.formatter.prettify_headers:
                        writer.writerow([row.get(key, "") for key in fieldnames])
                    else:
                        writer.writerow(row)

                # Flush periodically to free memory
                temp_file.flush()

            temp_file.close()
            return self._build_filename(filename), temp_file.name

        except Exception:
            temp_file.close()
            if os.path.exists(temp_file.name):
                os.unlink(temp_file.name)
            raise

    def _export_chunked_json_to_file(
        self, filename: str, queryset, chunk_size: int
    ) -> Tuple[str, str]:
        """Export to JSON in chunks, writing directly to a temp file."""
        temp_file = tempfile.NamedTemporaryFile(
            mode='w', delete=False, suffix='.json', encoding='utf-8'
        )
        try:
            temp_file.write("[")
            first_item = True

            for serialized_chunk in self._iter_chunks(queryset, chunk_size):
                for item in serialized_chunk:
                    if not first_item:
                        temp_file.write(",")
                    temp_file.write("\n  ")
                    temp_file.write(json.dumps(item, indent=None, default=str))
                    first_item = False

                # Flush periodically to free memory
                temp_file.flush()

            temp_file.write("\n]")
            temp_file.close()
            return self._build_filename(filename), temp_file.name

        except Exception:
            temp_file.close()
            if os.path.exists(temp_file.name):
                os.unlink(temp_file.name)
            raise

    def _export_chunked_xlsx_to_file(
        self, filename: str, queryset, chunk_size: int
    ) -> Tuple[str, str]:
        """Export to XLSX in chunks, writing directly to a temp file."""
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
        temp_file.close()  # Close so openpyxl can write to it

        try:
            # Use write_only mode for memory efficiency
            wb = Workbook(write_only=True)
            ws = wb.create_sheet()
            fieldnames = None

            for serialized_chunk in self._iter_chunks(queryset, chunk_size):
                # Initialize headers from first chunk
                if fieldnames is None:
                    fieldnames = self._extract_fieldnames(serialized_chunk)

                    # Write header row
                    if isinstance(self.formatter, XLSXFormatter) and self.formatter.prettify_headers:
                        headers = [self.formatter._prettify_header(key) for key in fieldnames]
                    else:
                        headers = fieldnames
                    ws.append(headers)

                # Write data rows
                for row in serialized_chunk:
                    formatted_row = []
                    for key in fieldnames:
                        value = row.get(key, "")
                        if isinstance(self.formatter, XLSXFormatter):
                            value = self.formatter._format_value(value)
                        formatted_row.append(value)
                    ws.append(formatted_row)

            wb.save(temp_file.name)
            return self._build_filename(filename), temp_file.name

        except Exception:
            if os.path.exists(temp_file.name):
                os.unlink(temp_file.name)
            raise
>>>>>>> Stashed changes
